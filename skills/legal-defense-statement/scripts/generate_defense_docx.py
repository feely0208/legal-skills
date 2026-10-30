#!/usr/bin/env python3
"""答辩状生成器 —— 将JSON答辩结果转为专业排版的Word文档(.docx)。"""

import json
import argparse
import os
from datetime import datetime

from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

# 色彩定义
COLORS = {
    "dark_blue": RGBColor(0x2F, 0x54, 0x96),
    "red":       RGBColor(0xDC, 0x35, 0x45),
    "green":     RGBColor(0x19, 0x87, 0x54),
    "orange":    RGBColor(0xFD, 0x7E, 0x14),
    "blue":      RGBColor(0x0D, 0x6E, 0xFD),
    "gray":      RGBColor(0x66, 0x66, 0x66),
    "light_gray":RGBColor(0x99, 0x99, 0x99),
}

POSITION_COLORS = {
    "驳回":         COLORS["red"],
    "部分认可":      COLORS["orange"],
    "认可但需澄清":  COLORS["blue"],
    "认可":          COLORS["green"],
}

FACT_POSITION_COLORS = {
    "属实":     COLORS["green"],
    "不属实":   COLORS["red"],
    "部分属实": COLORS["orange"],
}


def set_cell_shading(cell, color_hex):
    """设置Word表格单元格背景色"""
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading)


def add_run_to_para(para, text, font_name_en="Times New Roman",
                    font_name_cn="宋体", font_size=10.5,
                    bold=False, color=None):
    """向段落添加一个run，正确设置中英文字体"""
    run = para.add_run(text)
    run.font.size = Pt(font_size)
    run.font.name = font_name_en
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name_cn)
    run.bold = bold
    if color:
        run.font.color.rgb = color
    return run


def add_para(doc, text, color=None, bold=False, font_size=10.5,
             font_cn="宋体", font_en="Times New Roman",
             alignment=None, space_after=6, space_before=0,
             indent_first_line=None):
    """添加段落"""
    para = doc.add_paragraph()
    if alignment is not None:
        para.alignment = alignment
    pf = para.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    if indent_first_line:
        pf.first_line_indent = Cm(indent_first_line)
    add_run_to_para(para, text, font_name_en=font_en, font_name_cn=font_cn,
                    font_size=font_size, bold=bold, color=color)
    return para


def add_heading_cn(doc, text, level=1):
    """添加中文标题（黑体）"""
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = "黑体"
        run._element.rPr.rFonts.set(qn('w:eastAsia'), "黑体")
    return h


def format_party_info(data, party_type):
    """格式化当事人信息"""
    party = data.get(party_type, {})
    if not party:
        return ""

    lines = []
    if party.get("type") == "company":
        lines.append(party.get("name", ""))
        if party.get("credit_code"):
            lines.append(f'统一社会信用代码：{party["credit_code"]}')
        if party.get("address"):
            lines.append(f'住所地：{party["address"]}')
        if party.get("legal_rep"):
            lines.append(f'法定代表人：{party["legal_rep"]}，职务：{party.get("title", "")}')
        if party.get("phone"):
            lines.append(f'联系电话：{party["phone"]}')
    else:
        info_parts = [party.get("name", "")]
        if party.get("gender"):
            info_parts.append(party["gender"])
        if party.get("birth_date"):
            info_parts.append(f'{party["birth_date"]}出生')
        if party.get("ethnicity"):
            info_parts.append(party["ethnicity"])
        lines.append("，".join(info_parts))
        if party.get("id_number"):
            lines.append(f'身份证号码：{party["id_number"]}')
        if party.get("occupation"):
            lines.append(f'职业：{party["occupation"]}')
        if party.get("address"):
            lines.append(f'住址：{party["address"]}')
        if party.get("phone"):
            lines.append(f'联系电话：{party["phone"]}')
    return "\n".join(lines)


def render_defense_position(position):
    """渲染答辩态度，带颜色"""
    color = POSITION_COLORS.get(position, COLORS["gray"])
    return (position, color)


def render_fact_position(position):
    """渲染事实回应态度"""
    color = FACT_POSITION_COLORS.get(position, COLORS["gray"])
    return (position, color)


def build_claims_summary_table(doc, claims_response):
    """生成原告诉讼请求与答辩意见汇总表"""
    add_heading_cn(doc, "原告诉讼请求与答辩意见汇总", level=2)

    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    headers = ["序号", "原告诉讼请求", "答辩态度", "答辩理由概要"]
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        for p in hdr_cells[i].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.size = Pt(9)
                r.font.bold = True
                r.font.name = "黑体"
                r._element.rPr.rFonts.set(qn('w:eastAsia'), "黑体")
        set_cell_shading(hdr_cells[i], "2F5496")
        for p in hdr_cells[i].paragraphs:
            for r in p.runs:
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for item in claims_response:
        row = table.add_row()
        cells = row.cells
        cells[0].text = str(item.get("id", ""))
        cells[1].text = item.get("plaintiff_claim", "")
        pos, _ = render_defense_position(item.get("defense_position", ""))
        cells[2].text = pos
        cells[3].text = item.get("defense_reason", "")[:120]

        for cell in cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(8)
                    r.font.name = "宋体"
                    r._element.rPr.rFonts.set(qn('w:eastAsia'), "宋体")

        # 答辩态度列着色
        pos = item.get("defense_position", "")
        if pos == "驳回":
            set_cell_shading(cells[2], "FFCCCC")
        elif pos == "认可":
            set_cell_shading(cells[2], "D4EDDA")
        elif pos == "部分认可":
            set_cell_shading(cells[2], "FFF3CD")
        elif pos == "认可但需澄清":
            set_cell_shading(cells[2], "CCE5FF")

    widths = [Cm(1.0), Cm(4.0), Cm(2.0), Cm(5.0)]
    for row in table.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = w


def generate_report(data, output_path):
    """主函数：生成答辩状Word文档"""
    doc = Document()

    # 页面设置 (A4)
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.54)
    sec.bottom_margin = Cm(2.54)
    sec.left_margin = Cm(3.18)
    sec.right_margin = Cm(3.18)

    respondent_role = data.get("respondent_role", "defendant")
    role_label = "被告" if respondent_role == "defendant" else "原告"
    respondent = data.get("respondent", {})
    counterparty = data.get("counterparty", {})
    agent = data.get("agent", {})

    # ====== 封面 ======
    for _ in range(6):
        doc.add_paragraph()

    # 标题
    title = data.get("title", "民事答辩状")
    if "民事" in title:
        main_title = "民 事 答 辩 状"
    elif "行政" in title:
        main_title = "行 政 答 辩 状"
    else:
        main_title = title

    add_para(doc, main_title, bold=True, font_size=28,
             font_cn="黑体", alignment=WD_ALIGN_PARAGRAPH.CENTER,
             space_after=30)

    add_para(doc, title, font_size=16, font_cn="黑体",
             alignment=WD_ALIGN_PARAGRAPH.CENTER, space_after=20,
             color=RGBColor(0x33, 0x33, 0x33))

    info = [
        f'案号：{data.get("case_number", "—")}',
        f'审理法院：{data.get("court", "—")}',
        f'案由：{data.get("case_brief", "—")}',
        f'答辩人：{respondent.get("name", "—")}（{role_label}）',
        f'被答辩人：{counterparty.get("name", "—")}',
        f'提交日期：{data.get("submission_date", datetime.now().strftime("%Y年%m月%d日"))}',
    ]
    for line in info:
        add_para(doc, line, font_size=13,
                 alignment=WD_ALIGN_PARAGRAPH.CENTER, space_after=6,
                 color=RGBColor(0x55, 0x55, 0x55))

    doc.add_page_break()

    # ====== 正文开始 ======

    # 一、当事人信息
    add_heading_cn(doc, "一、当事人信息", level=1)

    add_para(doc, "答辩人：", bold=True, font_size=11, font_cn="黑体",
             space_before=8)
    respondent_info = format_party_info(data, "respondent")
    add_para(doc, respondent_info, font_size=11, indent_first_line=0.74,
             space_after=4)

    add_para(doc, "", font_size=4)

    add_para(doc, "被答辩人：", bold=True, font_size=11, font_cn="黑体",
             space_before=4)
    counterparty_info = format_party_info(data, "counterparty")
    add_para(doc, counterparty_info, font_size=11, indent_first_line=0.74,
             space_after=4)

    # 代理人
    if agent.get("name"):
        add_para(doc, "", font_size=4)
        add_para(doc, "委托代理人：", bold=True, font_size=11, font_cn="黑体")
        agent_parts = [agent.get("name", "")]
        if agent.get("law_firm"):
            agent_parts.append(agent["law_firm"])
        if agent.get("license_no"):
            agent_parts.append(f'执业证号：{agent["license_no"]}')
        if agent.get("phone"):
            agent_parts.append(f'联系电话：{agent["phone"]}')
        add_para(doc, "，".join(agent_parts), font_size=11,
                 indent_first_line=0.74)

    # 案号、案由、引言
    add_para(doc, "", font_size=4, space_after=8)
    case_number = data.get("case_number", "—")
    case_brief = data.get("case_brief", "—")
    add_para(doc,
             f'答辩人因被答辩人诉答辩人{case_brief}一案（案号：{case_number}），现提出答辩意见如下：',
             font_size=11, indent_first_line=0.74, space_after=8, bold=True)

    # ====== 二、答辩请求 ======
    doc.add_page_break()
    add_heading_cn(doc, "二、答辩请求", level=1)

    claims_response = data.get("claims_response", [])
    for item in claims_response:
        idx = item.get("id", "")
        claim = item.get("plaintiff_claim", "")
        position = item.get("defense_position", "")
        detail = item.get("defense_detail", "")
        reason = item.get("defense_reason", "")

        pos_text, pos_color = render_defense_position(position)

        add_para(doc, f'{idx}. 针对被答辩人关于「{claim}」的诉讼请求：',
                 bold=True, font_size=11, font_cn="黑体", space_before=10,
                 space_after=2)

        if detail:
            add_para(doc, f'答辩态度：{pos_text}', bold=True, font_size=11,
                     color=pos_color, indent_first_line=0.74, space_after=2)

        if reason:
            add_para(doc, "答辩理由：", bold=True, font_size=10.5,
                     font_cn="黑体", indent_first_line=0.74, space_after=2)
            for para_text in reason.split("\n"):
                para_text = para_text.strip()
                if para_text:
                    add_para(doc, para_text, font_size=10.5,
                             indent_first_line=0.74, space_after=4,
                             color=RGBColor(0x33, 0x33, 0x33))

        add_para(doc, "", font_size=2, space_after=4)

    # ====== 三、事实与理由 ======
    doc.add_page_break()
    add_heading_cn(doc, "三、事实与理由", level=1)

    facts = data.get("facts_and_reasons", {})

    # （一）对原告所述事实的回应
    response_items = facts.get("response_to_plaintiff_facts", [])
    if response_items:
        add_para(doc, "（一）对被答辩人所述事实的回应", bold=True, font_size=13,
                 font_cn="黑体", space_before=8, space_after=8)

        for i, item in enumerate(response_items, 1):
            allegation = item.get("plaintiff_allegation", "")
            position = item.get("respondent_position", "")
            response = item.get("respondent_response", "")

            pos_text, pos_color = render_fact_position(position)

            add_para(doc, f'{i}. 关于被答辩人主张的「{allegation}」', bold=True,
                     font_size=11, font_cn="黑体", space_before=8, space_after=2)

            add_para(doc, f'答辩人认为：{pos_text}。', bold=True, font_size=10.5,
                     color=pos_color, indent_first_line=0.74, space_after=2)

            if response:
                add_para(doc, response, font_size=10.5, indent_first_line=0.74,
                         color=RGBColor(0x33, 0x33, 0x33), space_after=4)

    # （二）本案基本事实
    respondent_facts = facts.get("respondent_facts", "")
    if respondent_facts:
        add_para(doc, "（二）本案基本事实", bold=True, font_size=13,
                 font_cn="黑体", space_before=12, space_after=6)

        for line in respondent_facts.split("\n"):
            line = line.strip()
            if line:
                add_para(doc, line, font_size=10.5, indent_first_line=0.74,
                         space_after=4)

    # （三）法律适用意见
    legal_basis = facts.get("legal_basis", [])
    if legal_basis:
        add_para(doc, "（三）法律适用意见", bold=True, font_size=13,
                 font_cn="黑体", space_before=12, space_after=6)

        for i, item in enumerate(legal_basis, 1):
            law = item.get("law", "")
            application = item.get("application", "")
            add_para(doc, f'{i}. {law}', bold=True, font_size=10.5,
                     color=COLORS["blue"], indent_first_line=0.74, space_after=2)
            if application:
                add_para(doc, application, font_size=10.5, indent_first_line=0.74,
                         color=RGBColor(0x33, 0x33, 0x33), space_after=4)

    # （四）对原告证据体系的评价
    evidence_assessment = facts.get("evidence_assessment", "")
    if evidence_assessment:
        add_para(doc, "（四）对被答辩人证据体系的意见", bold=True, font_size=13,
                 font_cn="黑体", space_before=12, space_after=6)
        add_para(doc, evidence_assessment, font_size=10.5, indent_first_line=0.74,
                 color=RGBColor(0x55, 0x55, 0x55))

    # 程序性抗辩
    procedural = facts.get("procedural_defenses", [])
    if procedural:
        active_defenses = [d for d in procedural if d.get("applicable")]
        if active_defenses:
            add_para(doc, "（五）程序性抗辩", bold=True, font_size=13,
                     font_cn="黑体", space_before=12, space_after=6)
            for d in active_defenses:
                add_para(doc, f'● {d.get("type", "")}：{d.get("note", "")}',
                         font_size=10.5, indent_first_line=0.74,
                         color=COLORS["red"], space_after=4)

    # ====== 四、结论 ======
    doc.add_page_break()
    add_heading_cn(doc, "四、结论", level=1)

    conclusion = data.get("conclusion", "")
    if conclusion:
        add_para(doc, conclusion, font_size=11, indent_first_line=0.74,
                 space_after=12)

    add_para(doc, "综上所述，恳请人民法院查明事实，依法公正裁判。",
             font_size=11, indent_first_line=0.74, bold=True, space_after=16)

    # 此致 + 法院
    add_para(doc, "此致", font_size=11, font_cn="黑体",
             alignment=WD_ALIGN_PARAGRAPH.LEFT, space_after=4)
    add_para(doc, data.get("court", "XX人民法院"), font_size=12, bold=True,
             font_cn="黑体", alignment=WD_ALIGN_PARAGRAPH.LEFT, space_after=16)

    # 附件
    appendices = data.get("appendices", [])
    if appendices:
        add_para(doc, "附：", bold=True, font_size=10.5, font_cn="黑体",
                 space_after=4)
        for item in appendices:
            add_para(doc, item, font_size=10, indent_first_line=0.74,
                     space_after=2)
        add_para(doc, "", font_size=4, space_after=8)

    # 签名栏
    add_para(doc, f'答辩人（签名/盖章）：', font_size=11, font_cn="黑体",
             alignment=WD_ALIGN_PARAGRAPH.RIGHT)
    if agent.get("name"):
        add_para(doc, f'委托代理人（签名）：', font_size=11, font_cn="黑体",
                 alignment=WD_ALIGN_PARAGRAPH.RIGHT, space_after=8)
    else:
        add_para(doc, "", font_size=4, space_after=8)
    add_para(doc, f'日期：{" " * 20}', font_size=11, font_cn="黑体",
             alignment=WD_ALIGN_PARAGRAPH.RIGHT)

    # 页脚声明
    add_para(doc, "", font_size=6, space_after=12)
    add_para(doc, "本文件由AI法律文书系统生成，仅供律师参考，不构成正式法律意见。建议经执业律师复核后使用。",
             font_size=8, color=COLORS["light_gray"],
             alignment=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, f'生成时间：{datetime.now().strftime("%Y年%m月%d日 %H:%M")}',
             font_size=8, color=COLORS["light_gray"],
             alignment=WD_ALIGN_PARAGRAPH.CENTER)

    # 保存
    out_dir = os.path.dirname(output_path)
    if out_dir and not os.path.exists(out_dir):
        os.makedirs(out_dir, exist_ok=True)
    doc.save(output_path)
    return output_path


def main():
    parser = argparse.ArgumentParser(description="答辩状生成器")
    parser.add_argument("--input", "-i", required=True, help="答辩状JSON文件路径")
    parser.add_argument("--output", "-o", required=True, help="输出Word文档路径(.docx)")
    args = parser.parse_args()

    if not os.path.exists(args.input):
        print(f"错误：找不到输入文件 —— {args.input}")
        return 1

    try:
        with open(args.input, "r", encoding="utf-8") as f:
            data = json.load(f)
    except json.JSONDecodeError as e:
        print(f"JSON解析错误：{e}")
        print("提示：请检查JSON字符串值中是否含有未转义的中文引号。")
        print("解决方案：将JSON值内部的中文双引号替换为「」或『』。")
        return 1

    if "claims_response" not in data:
        print("错误：JSON缺少 'claims_response' 字段")
        return 1

    path = generate_report(data, args.output)
    print(f"答辩状已生成：{path}")
    print(f"共 {len(data.get('claims_response', []))} 项诉讼请求答辩意见")
    return 0


if __name__ == "__main__":
    exit(main())
