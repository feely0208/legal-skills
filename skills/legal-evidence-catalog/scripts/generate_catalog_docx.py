#!/usr/bin/env python3
"""证据目录生成器 —— 将JSON证据目录转为专业排版的Word文档(.docx)。"""

import json
import argparse
import os
from datetime import datetime

from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml


def set_cell_shading(cell, color_hex):
    """设置Word表格单元格背景色"""
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading)


def add_run(para, text, font_name_en="Times New Roman",
            font_name_cn="宋体", font_size=10.5,
            bold=False, color=None):
    """向段落添加一个run，正确设置中英文字体"""
    run = para.add_run(text)
    run.font.size = Pt(font_size)
    run.font.name = font_name_en
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name_cn)
    run.bold = bold
    if color:
        run.font.color.rgb = color
    return run


def add_para(doc, text, color=None, bold=False, font_size=10.5,
             font_cn="宋体", font_en="Times New Roman",
             alignment=None, space_after=6, space_before=0,
             indent_first_line=None):
    """添加段落"""
    para = doc.add_paragraph()
    if alignment is not None:
        para.alignment = alignment
    pf = para.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    if indent_first_line:
        pf.first_line_indent = Cm(indent_first_line)
    add_run(para, text, font_name_en=font_en, font_name_cn=font_cn,
            font_size=font_size, bold=bold, color=color)
    return para


def add_heading_cn(doc, text, level=1):
    """添加中文标题（黑体）"""
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = "Times New Roman"
        run._element.rPr.rFonts.set(qn('w:eastAsia'), "黑体")
    return h


def build_evidence_table(doc, evidence_items):
    """生成证据目录核心表格"""
    add_heading_cn(doc, "证据目录", level=2)

    table = doc.add_table(rows=1, cols=9)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    headers = ["序号", "证据名称", "证据类型", "证据来源", "页数", "份数", "原件/复印件", "证明对象/证明目的", "备注"]
    hdr_cells = table.rows[0].cells

    for i, h in enumerate(headers):
        hdr_cells[i].text = ""
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.font.size = Pt(9)
        r.font.bold = True
        r.font.name = "Times New Roman"
        r._element.rPr.rFonts.set(qn('w:eastAsia'), "黑体")
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_cell_shading(hdr_cells[i], "2F5496")

    for item in evidence_items:
        row = table.add_row()
        cells = row.cells

        is_orig = "原件" if item.get("is_original", False) else "复印件"

        values = [
            str(item.get("id", "")),
            item.get("evidence_name", ""),
            item.get("evidence_type", ""),
            item.get("evidence_source", ""),
            str(item.get("pages", "")),
            str(item.get("copies", "")),
            is_orig,
            item.get("proof_purpose", ""),
            item.get("remarks", ""),
        ]

        for i, v in enumerate(values):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            align = WD_ALIGN_PARAGRAPH.CENTER if i in (0, 2, 4, 5, 6) else WD_ALIGN_PARAGRAPH.LEFT
            p.alignment = align
            r = p.add_run(v)
            r.font.size = Pt(8.5)
            r.font.name = "Times New Roman"
            r._element.rPr.rFonts.set(qn('w:eastAsia'), "宋体")

        # Alternating row colors
        if item.get("id", 0) % 2 == 0:
            for cell in cells:
                set_cell_shading(cell, "F2F6FC")

    # Set column widths
    col_widths = [Cm(0.8), Cm(3.2), Cm(1.4), Cm(2.2), Cm(0.7), Cm(0.7), Cm(1.2), Cm(5.0), Cm(1.8)]
    for row in table.rows:
        for i, w in enumerate(col_widths):
            row.cells[i].width = w

    return table


def build_summary_section(doc, summary):
    """生成证据概况部分"""
    add_heading_cn(doc, "证据概况", level=2)

    add_para(doc,
             f'本证据目录共列证据 {summary.get("total_items", 0)} 项，合计 {summary.get("total_pages", 0)} 页。',
             bold=True, font_size=11, space_after=6)

    add_para(doc,
             f'其中原件 {summary.get("original_count", 0)} 份，复印件 {summary.get("copy_count", 0)} 份。',
             font_size=10.5, space_after=6)

    types = summary.get("evidence_types_used", [])
    if types:
        add_para(doc, f'证据类型包括：{"、".join(types)}。', font_size=10.5, space_after=6)

    facts = summary.get("main_facts", [])
    if facts:
        add_para(doc, "【主要证明事项】", bold=True, font_size=10.5, font_cn="黑体", space_before=4)
        for f in facts:
            add_para(doc, f"● {f}", font_size=10.5, indent_first_line=0.74)

    notes = summary.get("notes", "")
    if notes:
        add_para(doc, f"说明：{notes}", font_size=10, color=RGBColor(0x66, 0x66, 0x66), space_before=4)


def generate_catalog(data, output_path):
    """主函数"""
    doc = Document()

    # Page setup (A4 landscape for better table fit)
    sec = doc.sections[0]
    sec.page_width = Cm(29.7)
    sec.page_height = Cm(21.0)
    sec.orientation = 1  # landscape
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(1.5)
    sec.right_margin = Cm(1.5)

    # ====== Cover / Title area ======
    add_para(doc, "", font_size=6, space_after=8)

    add_para(doc, "证 据 目 录", bold=True, font_size=24,
             font_cn="黑体", alignment=WD_ALIGN_PARAGRAPH.CENTER,
             space_after=20)

    add_para(doc, data.get("title", "证据目录"),
             font_size=14, font_cn="黑体",
             alignment=WD_ALIGN_PARAGRAPH.CENTER, space_after=12,
             color=RGBColor(0x33, 0x33, 0x33))

    # Case info
    info_lines = [
        f'案　号：{data.get("case_number", "—")}',
        f'审理法院：{data.get("court", "—")}',
        f'提交人：{data.get("party_name", "—")}（{data.get("submitting_party", "—")}）',
    ]
    if data.get("agent_name"):
        info_lines.append(f'代理人：{data.get("agent_name", "")}')
    info_lines.append(f'提交日期：{data.get("submission_date", datetime.now().strftime("%Y年%m月%d日"))}')

    for line in info_lines:
        add_para(doc, line, font_size=11,
                 alignment=WD_ALIGN_PARAGRAPH.CENTER, space_after=4,
                 color=RGBColor(0x55, 0x55, 0x55))

    add_para(doc, "", font_size=6, space_after=8)

    # ====== Evidence summary section ======
    summary = data.get("summary", {})
    if summary:
        build_summary_section(doc, summary)

    add_para(doc, "", font_size=4)

    # ====== Evidence catalog table ======
    evidence_items = data.get("evidence_items", [])
    build_evidence_table(doc, evidence_items)

    # ====== Certification & Signature ======
    add_para(doc, "", font_size=8)

    certification = data.get("certification", "以上证据材料经本人核对，复印件与原件一致。")
    add_para(doc, "", font_size=4)
    add_para(doc, "提交人声明：", bold=True, font_size=11, font_cn="黑体", space_before=12)
    add_para(doc, certification, font_size=10.5, indent_first_line=0.74, space_after=16)

    add_para(doc, f'提交人（签名/盖章）：{" " * 16}', font_size=11, font_cn="黑体",
             alignment=WD_ALIGN_PARAGRAPH.RIGHT)
    if data.get("agent_name"):
        add_para(doc, f'代理人（签名）：{" " * 18}', font_size=11, font_cn="黑体",
                 alignment=WD_ALIGN_PARAGRAPH.RIGHT, space_after=6)
    add_para(doc, f'提交日期：{" " * 16}年{" " * 6}月{" " * 6}日', font_size=11, font_cn="黑体",
             alignment=WD_ALIGN_PARAGRAPH.RIGHT)

    # Footer
    add_para(doc, "", font_size=6, space_after=8)
    add_para(doc, "本文件由AI法律证据系统生成，仅供律师参考。建议经执业律师复核后使用。",
             font_size=8, color=RGBColor(0x99, 0x99, 0x99),
             alignment=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, f'生成时间：{datetime.now().strftime("%Y年%m月%d日 %H:%M")}',
             font_size=8, color=RGBColor(0x99, 0x99, 0x99),
             alignment=WD_ALIGN_PARAGRAPH.CENTER)

    # Save
    out_dir = os.path.dirname(output_path)
    if out_dir and not os.path.exists(out_dir):
        os.makedirs(out_dir, exist_ok=True)
    doc.save(output_path)
    return output_path


def main():
    parser = argparse.ArgumentParser(description="证据目录生成器")
    parser.add_argument("--input", "-i", required=True, help="证据目录JSON文件路径")
    parser.add_argument("--output", "-o", required=True, help="输出Word文档路径(.docx)")
    args = parser.parse_args()

    if not os.path.exists(args.input):
        print(f"错误：找不到输入文件 —— {args.input}")
        return 1

    try:
        with open(args.input, "r", encoding="utf-8") as f:
            data = json.load(f)
    except json.JSONDecodeError as e:
        print(f"JSON解析错误：{e}")
        print("提示：请检查JSON字符串值中是否含有未转义的中文引号。")
        print("解决方案：将JSON值内部的中文双引号替换为「」或『』。")
        return 1

    if "evidence_items" not in data:
        print("错误：JSON缺少 'evidence_items' 字段")
        return 1

    path = generate_catalog(data, args.output)
    print(f"证据目录已生成：{path}")
    print(f"共 {len(data.get('evidence_items', []))} 项证据")
    return 0


if __name__ == "__main__":
    exit(main())
