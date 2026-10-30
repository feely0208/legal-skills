#!/usr/bin/env python3
"""合同审阅报告生成器 —— 将JSON审阅结果转为四色标记的Word文档。"""

import json
import argparse
import os
from datetime import datetime

from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

# 四色定义
COLORS = {
    "red":    RGBColor(0xDC, 0x35, 0x45),
    "orange": RGBColor(0xFD, 0x7E, 0x14),
    "blue":   RGBColor(0x0D, 0x6E, 0xFD),
    "green":  RGBColor(0x19, 0x87, 0x54),
}

COLOR_LABELS = {
    "red": "漏洞/缺陷",
    "orange": "风险提醒",
    "green": "修改建议",
}


def set_cell_shading(cell, color_hex):
    """设置Word表格单元格背景色"""
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading)


def add_run_to_para(para, text, font_name_en="Times New Roman",
                    font_name_cn="宋体", font_size=10.5,
                    bold=False, color=None):
    """向段落添加一个run，正确设置中英文字体"""
    run = para.add_run(text)
    run.font.size = Pt(font_size)
    run.font.name = font_name_en
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name_cn)
    run.bold = bold
    if color:
        run.font.color.rgb = color
    return run


def add_para(doc, text, color=None, bold=False, font_size=10.5,
             font_cn="宋体", font_en="Times New Roman",
             alignment=None, space_after=6, space_before=0,
             indent_first_line=None):
    """添加段落"""
    para = doc.add_paragraph()
    if alignment is not None:
        para.alignment = alignment
    pf = para.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    if indent_first_line:
        pf.first_line_indent = Cm(indent_first_line)
    add_run_to_para(para, text, font_name_en=font_en, font_name_cn=font_cn,
                    font_size=font_size, bold=bold, color=color)
    return para


def add_heading_cn(doc, text, level=1):
    """添加中文标题（黑体）"""
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = "黑体"
        run._element.rPr.rFonts.set(qn('w:eastAsia'), "黑体")
    return h


def add_section_tag(doc, label, color):
    """添加彩色标签如【漏洞/缺陷】"""
    para = doc.add_paragraph()
    pf = para.paragraph_format
    pf.space_after = Pt(4)
    pf.space_before = Pt(8)
    r = para.add_run(f"【{label}】")
    r.font.size = Pt(10)
    r.font.bold = True
    r.font.color.rgb = color
    r.font.name = "黑体"
    r._element.rPr.rFonts.set(qn('w:eastAsia'), "黑体")
    return para


def add_finding_section(doc, finding):
    """渲染单条审阅意见"""
    idx = finding.get("id", "")
    title = finding.get("title", "")
    cat = finding.get("category", "green")
    color = COLORS.get(cat, COLORS["green"])
    label = COLOR_LABELS.get(cat, "修改建议")

    # 标题行
    add_para(doc, f'第{idx}条  {title}', color=color, bold=True,
             font_size=12, font_cn="黑体", space_before=14)

    # 分类标签
    add_section_tag(doc, label, color)

    # 涉及条款
    add_para(doc, f'涉及条款：{finding.get("clause", "—")}', font_size=10.5)

    # 原条款
    txt = finding.get("original_text", "")
    if txt:
        add_para(doc, "【原条款内容】", bold=True, font_size=10.5,
                 font_cn="黑体", space_before=4)
        add_para(doc, txt, font_size=10.5, indent_first_line=0.74)

    # 问题分析
    txt = finding.get("description", "")
    if txt:
        add_para(doc, "【问题分析】", bold=True, font_size=10.5,
                 font_cn="黑体", space_before=4)
        add_para(doc, txt, color=color, font_size=10.5, indent_first_line=0.74)

    # 法律依据 — 蓝色
    txt = finding.get("legal_basis", "")
    if txt:
        add_para(doc, "【法律依据】", bold=True, font_size=10.5,
                 font_cn="黑体", color=COLORS["blue"], space_before=4)
        for line in txt.split("\n"):
            line = line.strip()
            if line:
                add_para(doc, line, color=COLORS["blue"], font_size=10.5,
                         indent_first_line=0.74)

    # 参考案例 — 蓝色
    txt = finding.get("case_reference", "")
    if txt:
        add_para(doc, "【参考案例】", bold=True, font_size=10.5,
                 font_cn="黑体", color=COLORS["blue"], space_before=4)
        for line in txt.split("\n"):
            line = line.strip()
            if line:
                add_para(doc, line, color=COLORS["blue"], font_size=10.5,
                         indent_first_line=0.74)

    # 修改建议 — 绿色
    txt = finding.get("suggestion", "")
    if txt:
        add_para(doc, "【修改建议】", bold=True, font_size=10.5,
                 font_cn="黑体", color=COLORS["green"], space_before=4)
        add_para(doc, txt, color=COLORS["green"], font_size=10.5,
                 indent_first_line=0.74)

    # 建议修改后文本 — 绿色
    txt = finding.get("revised_text", "")
    if txt:
        add_para(doc, "【建议修改后文本】", bold=True, font_size=10.5,
                 font_cn="黑体", color=COLORS["green"], space_before=4)
        add_para(doc, txt, color=COLORS["green"], font_size=10.5,
                 indent_first_line=0.74)

    # 风险等级 + 紧急程度
    add_para(doc,
             f'风险等级：{finding.get("risk_level", "—")}　|　紧急程度：{finding.get("urgency", "—")}',
             font_size=10, color=color, space_before=6)

    # 分隔线
    add_para(doc, "—" * 40, font_size=8,
             color=RGBColor(0xCC, 0xCC, 0xCC), space_before=4)


def build_summary_table(doc, findings):
    """生成风险汇总表格"""
    add_heading_cn(doc, "风险汇总表", level=2)

    table = doc.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    headers = ["序号", "类别", "涉及条款", "问题概述", "风险等级", "紧急程度"]
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        for p in hdr_cells[i].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.size = Pt(9)
                r.font.bold = True
                r.font.name = "黑体"
                r._element.rPr.rFonts.set(qn('w:eastAsia'), "黑体")
        set_cell_shading(hdr_cells[i], "2F5496")
        for p in hdr_cells[i].paragraphs:
            for r in p.runs:
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for f in findings:
        row = table.add_row()
        cells = row.cells
        cells[0].text = str(f.get("id", ""))
        cells[1].text = COLOR_LABELS.get(f.get("category", ""), "")
        cells[2].text = f.get("clause", "")
        cells[3].text = f.get("title", "")
        cells[4].text = f.get("risk_level", "")
        cells[5].text = f.get("urgency", "")

        for cell in cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
                    r.font.name = "宋体"
                    r._element.rPr.rFonts.set(qn('w:eastAsia'), "宋体")

        # 风险等级底色
        rl = f.get("risk_level", "")
        if rl == "高":
            set_cell_shading(cells[4], "FFCCCC")
        elif rl == "中":
            set_cell_shading(cells[4], "FFF3CD")
        elif rl == "低":
            set_cell_shading(cells[4], "D4EDDA")

        # 类别底色
        c = f.get("category", "")
        if c == "red":
            set_cell_shading(cells[1], "F8D7DA")
        elif c == "orange":
            set_cell_shading(cells[1], "FFE5CC")
        elif c == "green":
            set_cell_shading(cells[1], "D4EDDA")

    # 列宽
    widths = [Cm(1.0), Cm(1.8), Cm(2.5), Cm(5.5), Cm(1.5), Cm(1.8)]
    for row in table.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = w


def generate_report(data, output_path):
    """主函数"""
    doc = Document()

    # 页面设置 (A4)
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.54)
    sec.bottom_margin = Cm(2.54)
    sec.left_margin = Cm(3.18)
    sec.right_margin = Cm(3.18)

    # ====== 封面 ======
    for _ in range(6):
        doc.add_paragraph()

    add_para(doc, "合同审阅报告", bold=True, font_size=26,
             font_cn="黑体", alignment=WD_ALIGN_PARAGRAPH.CENTER,
             space_after=24)

    add_para(doc, data.get("title", "合同审阅报告"),
             font_size=16, font_cn="黑体",
             alignment=WD_ALIGN_PARAGRAPH.CENTER, space_after=12,
             color=RGBColor(0x33, 0x33, 0x33))

    info = [
        f'合同类型：{data.get("contract_type", "—")}',
        f'委托人：{data.get("client_name", "—")}',
        f'审阅日期：{data.get("review_date", datetime.now().strftime("%Y-%m-%d"))}',
        f'审阅人：{data.get("reviewer", "—")}',
    ]
    for line in info:
        add_para(doc, line, font_size=12,
                 alignment=WD_ALIGN_PARAGRAPH.CENTER, space_after=4,
                 color=RGBColor(0x55, 0x55, 0x55))

    doc.add_page_break()

    # ====== 摘要 ======
    s = data.get("summary", {})
    add_heading_cn(doc, "一、审阅摘要", level=1)

    overall = s.get("overall_risk", "—")
    signability = s.get("signability", "—")
    total = s.get("total_issues", 0)
    rc = s.get("red_count", 0)
    oc = s.get("orange_count", 0)
    gc = s.get("green_count", 0)

    add_para(doc,
             f'整体风险：{overall}　|　签约可行性：{signability}　|　共发现 {total} 个问题',
             bold=True, font_size=11, space_after=8)

    add_para(doc,
             f'■ 漏洞/缺陷 {rc} 项　■ 风险提醒 {oc} 项　■ 修改建议 {gc} 项',
             font_size=10.5, space_after=8, color=RGBColor(0x55, 0x55, 0x55))

    brief = s.get("brief", "")
    if brief:
        add_para(doc, "【整体评估】", bold=True, font_size=10.5, font_cn="黑体")
        add_para(doc, brief, font_size=10.5, indent_first_line=0.74)

    bl = s.get("bottom_lines", [])
    if bl:
        add_para(doc, "【谈判底线 / 必须修改】", bold=True,
                 font_size=10.5, font_cn="黑体", color=COLORS["red"],
                 space_before=6)
        for item in bl:
            add_para(doc, f"● {item}", font_size=10.5,
                     color=COLORS["red"], indent_first_line=0.74)

    ng = s.get("negotiable", [])
    if ng:
        add_para(doc, "【可协商条款】", bold=True,
                 font_size=10.5, font_cn="黑体", color=COLORS["orange"],
                 space_before=6)
        for item in ng:
            add_para(doc, f"● {item}", font_size=10.5,
                     color=COLORS["orange"], indent_first_line=0.74)

    mr = s.get("max_risk", "")
    if mr:
        add_para(doc, "【最不利风险敞口】", bold=True,
                 font_size=10.5, font_cn="黑体", color=COLORS["red"],
                 space_before=6)
        add_para(doc, mr, font_size=10.5, color=COLORS["red"],
                 indent_first_line=0.74)

    lt = s.get("litigation_trend", "")
    if lt:
        add_para(doc, "【类案裁判趋势】", bold=True,
                 font_size=10.5, font_cn="黑体", color=COLORS["blue"],
                 space_before=6)
        add_para(doc, lt, font_size=10.5, color=COLORS["blue"],
                 indent_first_line=0.74)

    # ====== 风险汇总表 ======
    doc.add_page_break()
    add_heading_cn(doc, "二、风险汇总表", level=1)
    findings = data.get("findings", [])
    build_summary_table(doc, findings)

    # ====== 逐条审阅 ======
    doc.add_page_break()
    add_heading_cn(doc, "三、逐条审阅意见", level=1)

    # 图例
    add_para(doc, "【颜色图例】", bold=True, font_size=10, font_cn="黑体")
    legend = [
        ("■ 红色 = 漏洞/缺陷（缺失必备条款、违反强制性规定、可能无效等）", COLORS["red"]),
        ("■ 橙色 = 风险提醒（对己方不利条款、权利义务失衡、高概率争议点）", COLORS["orange"]),
        ("■ 蓝色 = 法律依据（现行法条与裁判案例要旨）", COLORS["blue"]),
        ("■ 绿色 = 修改建议（具体修改方案、替代文本、谈判策略）", COLORS["green"]),
    ]
    for txt, c in legend:
        add_para(doc, txt, color=c, font_size=9, space_after=2)

    add_para(doc, "", font_size=6)

    for f in findings:
        add_finding_section(doc, f)

    # 页脚
    add_para(doc, "", font_size=6, space_after=12)
    add_para(doc, "本报告由AI法律审阅系统生成，仅供律师参考，不构成正式法律意见。",
             font_size=8, color=RGBColor(0x99, 0x99, 0x99),
             alignment=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, f'生成时间：{datetime.now().strftime("%Y年%m月%d日 %H:%M")}',
             font_size=8, color=RGBColor(0x99, 0x99, 0x99),
             alignment=WD_ALIGN_PARAGRAPH.CENTER)

    # 保存
    out_dir = os.path.dirname(output_path)
    if out_dir and not os.path.exists(out_dir):
        os.makedirs(out_dir, exist_ok=True)
    doc.save(output_path)
    return output_path


def main():
    parser = argparse.ArgumentParser(description="合同审阅报告生成器")
    parser.add_argument("--input", "-i", required=True, help="审阅JSON文件路径")
    parser.add_argument("--output", "-o", required=True, help="输出Word文档路径(.docx)")
    args = parser.parse_args()

    if not os.path.exists(args.input):
        print(f"错误：找不到输入文件 —— {args.input}")
        return 1

    try:
        with open(args.input, "r", encoding="utf-8") as f:
            data = json.load(f)
    except json.JSONDecodeError as e:
        print(f"JSON解析错误：{e}")
        print("提示：请检查JSON字符串值中是否含有未转义的中文引号「\"\"」。")
        print("解决方案：将JSON值内部的中文双引号替换为「」或『』。")
        return 1

    if "findings" not in data:
        print("错误：JSON缺少 'findings' 字段")
        return 1

    path = generate_report(data, args.output)
    print(f"审阅报告已生成：{path}")
    print(f"共 {len(data.get('findings', []))} 条审阅意见")
    return 0


if __name__ == "__main__":
    exit(main())
