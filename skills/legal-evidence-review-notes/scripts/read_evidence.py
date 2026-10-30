#!/usr/bin/env python3
"""多格式案卷文件读取器。支持 PDF/DOCX/DOC/XLSX/XLS/JPG/JPEG/PNG/BMP/CSV/TXT/JSON/HTML。"""

import os
import sys
import argparse


def read_txt(path):
    """读取纯文本文件，自动检测编码"""
    for enc in ['utf-8', 'utf-16', 'gbk', 'gb18030', 'latin-1']:
        try:
            with open(path, 'r', encoding=enc) as f:
                return f.read()
        except (UnicodeDecodeError, UnicodeError):
            continue
    with open(path, 'r', encoding='utf-8', errors='replace') as f:
        return f.read()


def read_pdf(path):
    """读取PDF文件，优先使用pdfplumber，回退到PyPDF2"""
    try:
        import pdfplumber
        text_parts = []
        with pdfplumber.open(path) as pdf:
            for page in pdf.pages:
                t = page.extract_text()
                if t:
                    text_parts.append(t)
        return '\n'.join(text_parts)
    except ImportError:
        pass
    try:
        from PyPDF2 import PdfReader
        reader = PdfReader(path)
        return '\n'.join(p.extract_text() or '' for p in reader.pages)
    except ImportError:
        raise ImportError("需要安装 pdfplumber 或 PyPDF2: pip install pdfplumber")


def read_docx(path):
    """读取DOCX文件"""
    from docx import Document
    doc = Document(path)
    parts = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text for cell in row.cells]
            parts.append(' | '.join(cells))
    return '\n'.join(parts)


def read_doc(path):
    """读取旧版DOC文件 (OLE2格式)"""
    import olefile
    ole = olefile.OleFileIO(path)
    if ole.exists('WordDocument'):
        stream = ole.openstream('WordDocument')
        data = stream.read()
        import re
        text = data.decode('utf-8', errors='replace')
        text = re.sub(r'[^\x20-\x7e一-鿿　-〿＀-￯\n\r]', '', text)
        ole.close()
        return text
    ole.close()
    return read_txt(path)


def read_xlsx(path):
    """读取XLSX文件"""
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    parts = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        parts.append(f'--- Sheet: {sheet_name} ---')
        for row in ws.iter_rows(values_only=True):
            row_text = ' | '.join(str(c) if c is not None else '' for c in row)
            if row_text.strip().replace('|', '').strip():
                parts.append(row_text)
    return '\n'.join(parts)


def read_xls(path):
    """读取旧版XLS文件"""
    import xlrd
    wb = xlrd.open_workbook(path)
    parts = []
    for sheet in wb.sheets():
        parts.append(f'--- Sheet: {sheet.name} ---')
        for row_idx in range(sheet.nrows):
            row = sheet.row_values(row_idx)
            row_text = ' | '.join(str(c) if c else '' for c in row)
            if row_text.strip().replace('|', '').strip():
                parts.append(row_text)
    return '\n'.join(parts)


def read_csv(path):
    """读取CSV文件"""
    import csv
    parts = []
    for enc in ['utf-8', 'utf-8-sig', 'gbk', 'gb18030']:
        try:
            with open(path, 'r', encoding=enc, newline='') as f:
                reader = csv.reader(f)
                for row in reader:
                    parts.append(' | '.join(row))
            return '\n'.join(parts)
        except (UnicodeDecodeError, UnicodeError):
            continue
    return read_txt(path)


def read_json(path):
    """读取JSON文件，格式化为可读文本"""
    import json
    with open(path, 'r', encoding='utf-8') as f:
        data = json.load(f)
    return json.dumps(data, ensure_ascii=False, indent=2)


def read_html(path):
    """读取HTML文件，提取纯文本"""
    from bs4 import BeautifulSoup
    with open(path, 'r', encoding='utf-8') as f:
        content = f.read()
    soup = BeautifulSoup(content, 'html.parser')
    for tag in soup(['script', 'style', 'meta', 'link']):
        tag.decompose()
    return soup.get_text('\n', strip=True)


def read_image(path):
    """读取图片文件 —— 返回图片路径，由AI直接读取识别内容"""
    return f"[IMAGE]{path}"


READERS = {
    '.pdf':  read_pdf,
    '.docx': read_docx,
    '.doc':  read_doc,
    '.xlsx': read_xlsx,
    '.xls':  read_xls,
    '.csv':  read_csv,
    '.txt':  read_txt,
    '.json': read_json,
    '.html': read_html,
    '.htm':  read_html,
    '.jpg':  read_image,
    '.jpeg': read_image,
    '.png':  read_image,
    '.bmp':  read_image,
}


def read_evidence(path):
    """读取任意支持格式的案卷文件"""
    ext = os.path.splitext(path)[1].lower()
    if ext not in READERS:
        supported = ', '.join(READERS.keys())
        raise ValueError(f"不支持的格式: {ext}\n支持的格式: {supported}")
    return READERS[ext](path)


def main():
    parser = argparse.ArgumentParser(description='多格式案卷文件读取器')
    parser.add_argument('input', help='案卷文件路径')
    parser.add_argument('--output', '-o', default=None, help='输出文本文件路径（可选）')
    parser.add_argument('--max-chars', '-m', type=int, default=50000,
                        help='最大输出字符数（默认50000）')
    args = parser.parse_args()

    if not os.path.exists(args.input):
        print(f"错误：文件不存在 —— {args.input}", file=sys.stderr)
        return 1

    try:
        result = read_evidence(args.input)
    except Exception as e:
        print(f"读取失败: {e}", file=sys.stderr)
        return 1

    ext = os.path.splitext(args.input)[1].lower()

    if ext in ('.jpg', '.jpeg', '.png', '.bmp'):
        print(f"格式: {ext} (图片)")
        print(f"图片路径: {result}")
    else:
        print(f"格式: {ext}")
        print(f"总字符数: {len(result)}")

        if len(result) > args.max_chars:
            result = result[:args.max_chars]
            print(f"注意: 已截断至 {args.max_chars} 字符")

        if args.output:
            os.makedirs(os.path.dirname(args.output) or '.', exist_ok=True)
            with open(args.output, 'w', encoding='utf-8') as f:
                f.write(result)
            print(f"已保存到: {args.output}")
        else:
            print(result)

    return 0


if __name__ == '__main__':
    exit(main())
